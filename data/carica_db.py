#!/usr/bin/env python3
"""Carica in data/ordini.db i PDF di data/pdf/ non ancora presenti.

    python3 data/carica_db.py                 # carica i PDF nuovi
    python3 data/carica_db.py --riclassifica  # riapplica categorie.json a tutte le righe
    python3 data/carica_db.py --report        # solo tabelle di controllo
"""
import argparse
import csv
import datetime
import json
import os
import re
import sqlite3
import sys

QUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, QUI)
from parse_pdf import parse_pdf  # noqa: E402

PDF_DIR = os.path.join(QUI, "pdf")
LOG_DIR = os.path.join(QUI, "log")
DB = os.path.join(QUI, "ordini.db")
ANAGRAFICA = os.path.normpath(os.path.join(QUI, "..", "..", "MAK 2026.xlsx"))
TOLLERANZA = 0.05


def apri():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    with open(os.path.join(QUI, "schema.sql"), encoding="utf-8") as f:
        con.executescript(f.read())
    # DB creato prima dell'aggiunta di network/tipo: si aggiungono le colonne
    colonne = {r["name"] for r in con.execute("PRAGMA table_info(clienti)")}
    for c in ("network", "tipo"):
        if c not in colonne:
            con.execute(f"ALTER TABLE clienti ADD COLUMN {c} TEXT")
    return con


def applica_profili(con):
    """Network e Tipo cliente dal corpo delle notifiche (data/log/profili.json):
    per ogni cliente vale l'ordine piu' recente che li riporta."""
    try:
        with open(os.path.join(LOG_DIR, "profili.json"), encoding="utf-8") as f:
            profili = json.load(f)
    except (OSError, ValueError):
        return 0
    visti = set()
    with con:
        for r in con.execute("SELECT id, codice_cliente FROM ordini ORDER BY data_ora DESC").fetchall():
            p = profili.get(r["id"])
            if p is None or r["codice_cliente"] in visti:
                continue
            visti.add(r["codice_cliente"])
            con.execute("UPDATE clienti SET network=?, tipo=? WHERE codice=?",
                        (p.get("network") or None, p.get("tipo") or None, r["codice_cliente"]))
    return len(visti)


def carica_regole():
    with open(os.path.join(QUI, "categorie.json"), encoding="utf-8") as f:
        c = json.load(f)
    for r in c["regole"]:
        if "regex_descrizione" in r:
            r["_re"] = re.compile(r["regex_descrizione"])
        if "regex_codice" in r:
            r["_rec"] = re.compile(r["regex_codice"])
    return c


def categoria(cod, desc, cfg):
    cod, desc = (cod or "").upper(), (desc or "").upper()
    for r in cfg["regole"]:
        if "prefisso" in r and not any(cod.startswith(p) for p in r["prefisso"]):
            continue
        if "contiene" in r and not any(k in desc for k in r["contiene"]):
            continue
        if "_re" in r and not r["_re"].search(desc):
            continue
        if "_rec" in r and not r["_rec"].search(cod):
            continue
        return r["categoria"]
    return cfg["default"]


def inserisci(con, o, cfg, nome_pdf):
    merce_righe = round(sum(r["importo"] for r in o["righe"]), 2)
    atteso = None
    quadra = None
    if o["imponibile"] is not None:
        atteso = round(o["imponibile"] - (o["spedizione"] or 0) - (o["pfu"] or 0)
                       - (o["contrassegno"] or 0) - (o["altri_costi"] or 0), 2)
        quadra = 1 if abs(merce_righe - atteso) <= TOLLERANZA else 0
    c = con.execute("SELECT ultimo_ordine FROM clienti WHERE codice=?", (o["codice_cliente"],)).fetchone()
    dati = (o["ragione_sociale"], o["citta"], o["provincia"], o["cap"], o["indirizzo"], o["piva"],
            o["cod_fiscale"], o["cod_network"], o["flag_cliente"])
    if c is None:
        con.execute("INSERT INTO clienti (codice, ragione_sociale, citta, provincia, cap, indirizzo, piva,"
                    " cod_fiscale, cod_network, flag_cliente) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (o["codice_cliente"],) + dati)
    elif (c["ultimo_ordine"] or "") <= (o["data_ora"] or ""):
        # i dati cliente li detta l'ordine piu' recente
        con.execute("UPDATE clienti SET ragione_sociale=?, citta=?, provincia=?, cap=?, indirizzo=?, piva=?,"
                    " cod_fiscale=?, cod_network=?, flag_cliente=? WHERE codice=?", dati + (o["codice_cliente"],))
    con.execute(
        "INSERT INTO ordini (id, codice_cliente, data_ora, inserito_da, utente_raw, kit_completo,"
        " termini_pagamento, vettore, nota, omaggio, merce_righe, merce, spedizione, pfu, contrassegno, altri_costi,"
        " imponibile, iva, totale, quadra, pagine, pdf) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (o["id"], o["codice_cliente"], o["data_ora"], o["inserito_da"], o["utente_raw"], o["kit_completo"],
         o["termini_pagamento"], o["vettore"], o.get("nota"),
         1 if "OMAGG" in (o["termini_pagamento"] or "").upper() else 0, merce_righe, o["merce"], o["spedizione"], o["pfu"],
         o["contrassegno"], o["altri_costi"] or None, o["imponibile"], o["iva"], o["totale"], quadra,
         o.get("pagine"), nome_pdf))
    for n, r in enumerate(o["righe"], 1):
        con.execute(
            "INSERT INTO righe (ordine_id, n_riga, cod_articolo, descrizione, applicazione, coupon, qta, prezzo,"
            " sconto, importo, categoria) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (o["id"], n, r["cod_articolo"], r["descrizione"], r["applicazione"], r.get("coupon"), r["qta"],
             r["prezzo"], r["sconto"], r["importo"], categoria(r["cod_articolo"], r["descrizione"], cfg)))
    con.execute("UPDATE clienti SET primo_ordine=(SELECT MIN(data_ora) FROM ordini WHERE codice_cliente=?),"
                " ultimo_ordine=(SELECT MAX(data_ora) FROM ordini WHERE codice_cliente=?) WHERE codice=?",
                (o["codice_cliente"],) * 3)


def carica(con, cfg, verbose=True):
    gia = {r["pdf"] for r in con.execute("SELECT pdf FROM ordini")}
    gia_id = {r["id"] for r in con.execute("SELECT id FROM ordini")}
    with con:   # i PDF scartati in passato si riprovano a ogni giro (il parser puo' essere migliorato)
        con.execute("DELETE FROM scarti WHERE origine='pdf'")
    nuovi = sorted(f for f in os.listdir(PDF_DIR) if f.endswith(".pdf") and f not in gia)
    ok = ko = 0
    oggi = datetime.datetime.now().isoformat(sep=" ", timespec="seconds")
    for k, f in enumerate(nuovi, 1):
        try:
            o = parse_pdf(os.path.join(PDF_DIR, f))
            if o["id"] != f[:-4]:
                # PDF salvato sotto il nome sbagliato (indice slittato in Mail): il contenuto fa fede.
                # Si rinomina; l'ordine atteso verra' riscaricato al prossimo giro.
                giusto = o["id"] + ".pdf"
                if o["id"] in gia_id or os.path.exists(os.path.join(PDF_DIR, giusto)):
                    os.remove(os.path.join(PDF_DIR, f))
                    continue
                os.rename(os.path.join(PDF_DIR, f), os.path.join(PDF_DIR, giusto))
                f = giusto
            if o["id"] in gia_id:
                raise ValueError(f"ordine {o['id']} gia' caricato da un altro PDF")
            with con:
                inserisci(con, o, cfg, f)
            gia_id.add(o["id"])
            ok += 1
        except Exception as e:
            ko += 1
            with con:
                con.execute("INSERT INTO scarti VALUES ('pdf', ?, ?, ?)", (f, str(e), oggi))
        if verbose and k % 500 == 0:
            print(f"  caricati {k}/{len(nuovi)}")
    return ok, ko


def riclassifica(con, cfg):
    with con:
        for r in con.execute("SELECT id, cod_articolo, descrizione FROM righe").fetchall():
            con.execute("UPDATE righe SET categoria=? WHERE id=?",
                        (categoria(r["cod_articolo"], r["descrizione"], cfg), r["id"]))


def segna_anagrafica(con):
    """Legge SOLO le colonne CLIENTE e COD. del foglio Clienti. Mai ID/PWD."""
    if not os.path.exists(ANAGRAFICA):
        return None
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(ANAGRAFICA, read_only=True, data_only=True)["Clienti"]
        intest = [c for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        i_cod = intest.index("COD.")
        codici = set()
        for (v,) in ws.iter_rows(min_row=2, min_col=i_cod + 1, max_col=i_cod + 1, values_only=True):
            if v is not None:
                codici.add(str(v).strip().zfill(6))
    with con:
        con.execute("UPDATE clienti SET in_anagrafica=0")
        con.executemany("UPDATE clienti SET in_anagrafica=1 WHERE codice=?", [(c,) for c in codici])
    return len(codici)


def report(con, verbose=True):
    os.makedirs(LOG_DIR, exist_ok=True)
    q = lambda s, *a: con.execute(s, a).fetchall()
    tot_merce = q("SELECT COALESCE(SUM(importo),0) s FROM righe")[0]["s"] or 0

    # frequenza per prefisso codice x categoria
    freq = {}
    for r in q("SELECT cod_articolo, descrizione, qta, importo, categoria FROM righe"):
        k = ((r["cod_articolo"] or "")[:4], r["categoria"])
        d = freq.setdefault(k, {"righe": 0, "qta": 0.0, "importo": 0.0, "esempi": []})
        d["righe"] += 1
        d["qta"] += r["qta"] or 0
        d["importo"] += r["importo"] or 0
        if len(d["esempi"]) < 3 and r["descrizione"] not in d["esempi"]:
            d["esempi"].append(r["descrizione"])
    with open(os.path.join(LOG_DIR, "categorie_frequenza.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["prefisso", "categoria", "righe", "qta", "importo", "esempio1", "esempio2", "esempio3"])
        for (pre, cat), d in sorted(freq.items(), key=lambda x: -x[1]["importo"]):
            w.writerow([pre, cat, d["righe"], round(d["qta"]), round(d["importo"], 2)] + d["esempi"])

    r = {
        "ordini": q("SELECT COUNT(*) n FROM ordini")[0]["n"],
        "clienti": q("SELECT COUNT(*) n FROM clienti")[0]["n"],
        "righe": q("SELECT COUNT(*) n FROM righe")[0]["n"],
        "merce": round(tot_merce, 2),
        "periodo": tuple(q("SELECT MIN(data_ora) a, MAX(data_ora) b FROM ordini")[0]),
        "omaggi": tuple(q("SELECT COUNT(*) n, COALESCE(SUM(merce_righe),0) s FROM ordini WHERE omaggio=1")[0]),
        "non_quadrano": [x["id"] for x in q("SELECT id FROM ordini WHERE quadra=0 ORDER BY id")],
        "senza_imponibile": q("SELECT COUNT(*) n FROM ordini WHERE quadra IS NULL")[0]["n"],
        "multipagina": [x["id"] for x in q("SELECT id FROM ordini WHERE pagine>1 ORDER BY id")],
        "scarti_pdf": q("SELECT COUNT(*) n FROM scarti WHERE origine='pdf'")[0]["n"],
        "fuori_anagrafica": q("SELECT COUNT(*) n FROM clienti WHERE in_anagrafica=0")[0]["n"],
        "senza_profilo": q("SELECT COUNT(*) n FROM clienti WHERE tipo IS NULL AND network IS NULL")[0]["n"],
        "categorie": [(x["categoria"], x["n"], round(x["s"] or 0, 2)) for x in q(
            "SELECT categoria, COUNT(*) n, SUM(importo) s FROM righe GROUP BY categoria ORDER BY s DESC")],
    }
    altro = next((s for c, _, s in r["categorie"] if c == "Altro"), 0.0)
    r["quota_altro"] = round(100 * altro / tot_merce, 2) if tot_merce else 0.0
    if verbose:
        print(f"Ordini {r['ordini']} | clienti {r['clienti']} | righe {r['righe']} | merce {r['merce']:.2f}")
        print(f"Periodo {r['periodo'][0]} -> {r['periodo'][1]}")
        print(f"Non quadrano: {len(r['non_quadrano'])} | senza imponibile: {r['senza_imponibile']} | "
              f"multipagina: {len(r['multipagina'])} | PDF scartati: {r['scarti_pdf']} | "
              f"clienti fuori anagrafica: {r['fuori_anagrafica']} | "
              f"omaggi: {r['omaggi'][0]} ({r['omaggi'][1]:.2f} a listino, esclusi dai valori in pagina)")
        for c, n, s in r["categorie"]:
            print(f"  {c:26s} {n:6d} righe  {s:12.2f}  {100 * s / tot_merce if tot_merce else 0:5.1f}%")
    return r


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--riclassifica", action="store_true")
    ap.add_argument("--report", action="store_true")
    a = ap.parse_args()
    con, cfg = apri(), carica_regole()
    if a.riclassifica:
        riclassifica(con, cfg)
    elif not a.report:
        ok, ko = carica(con, cfg)
        print(f"PDF caricati: {ok}, scartati: {ko}")
        segna_anagrafica(con)
    applica_profili(con)
    report(con)
    return 0


if __name__ == "__main__":
    sys.exit(main())
